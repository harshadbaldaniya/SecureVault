"""Safe parsing and validation of employee spreadsheets (untrusted input).

Defences, in order:
  1. size limit, ".xlsx" only
  2. ZIP inspection before any XML is parsed: valid OOXML package, no macros, bounded entry count and
     total decompressed size (zip-bomb guard)
  3. openpyxl in read-only, values-only mode; defusedxml is installed so openpyxl's XML parsing is hardened
     against entity-expansion and external-entity attacks
  4. row limit; every cell coerced to bounded plain text; formula-like text rejected (it would become a
     live formula when exported to a spreadsheet later)
"""

import io
import zipfile
from dataclasses import dataclass, field

from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.departments.models import Department

REQUIRED_COLUMNS = ("email", "full_name")
OPTIONAL_COLUMNS = ("department", "role")
HEADER_ALIASES = {
    "email": "email",
    "e-mail": "email",
    "email address": "email",
    "work email": "email",
    "full_name": "full_name",
    "full name": "full_name",
    "name": "full_name",
    "department": "department",
    "department code": "department",
    "dept": "department",
    "role": "role",
}
IMPORTABLE_ROLES = {"", "EMPLOYEE"}  # privileged roles are never granted by import
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
MAX_ZIP_ENTRIES = 1_000
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


class FileRejected(Exception):
    """The upload as a whole cannot be used. The message is safe to show to the uploader."""


@dataclass
class ParsedRow:
    row: int
    email: str
    full_name: str
    department: str
    role: str
    errors: list[dict] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {
            "row": self.row,
            "email": self.email,
            "full_name": self.full_name,
            "department": self.department,
            "role": self.role,
            "errors": self.errors,
        }


def check_upload(name: str, data: bytes) -> None:
    if not name.lower().endswith(".xlsx"):
        raise FileRejected("Upload an Excel workbook (.xlsx). Older .xls, .xlsm (macros) and CSV are not accepted.")
    if len(data) > settings.IMPORT_MAX_BYTES:
        raise FileRejected(f"The file is larger than {settings.IMPORT_MAX_BYTES // (1024 * 1024)} MB.")
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise FileRejected("This is not a valid .xlsx workbook.") from None
    with archive:
        entries = archive.infolist()
        names = {e.filename for e in entries}
        if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
            raise FileRejected("This is not a valid .xlsx workbook.")
        if any(n.lower().endswith("vbaproject.bin") for n in names):
            raise FileRejected("Workbooks containing macros are not accepted.")
        if len(entries) > MAX_ZIP_ENTRIES or sum(e.file_size for e in entries) > MAX_UNCOMPRESSED_BYTES:
            raise FileRejected("The workbook is too large when uncompressed.")


def _text(value, limit: int) -> str:
    if value is None:
        return ""
    return str(value).strip()[:limit]


def parse_rows(data: bytes) -> list[ParsedRow]:
    """Read the first worksheet. Raises FileRejected for file-level problems."""
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise FileRejected("The workbook could not be read.") from None
    try:
        return _read_sheet(workbook)
    except FileRejected:
        raise
    except Exception:
        # Read-only mode parses worksheet XML lazily while iterating, so malformed or hostile XML (e.g. entity
        # expansion stopped by defusedxml) surfaces here rather than in load_workbook.
        raise FileRejected("The workbook could not be read.") from None
    finally:
        workbook.close()


def _read_sheet(workbook) -> list[ParsedRow]:
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        raise FileRejected("The first worksheet is empty.")
    columns: dict[str, int] = {}
    for index, title in enumerate(header):
        key = HEADER_ALIASES.get(_text(title, 60).lower())
        if key and key not in columns:
            columns[key] = index
    missing = [c for c in REQUIRED_COLUMNS if c not in columns]
    if missing:
        raise FileRejected(f"Missing required column(s): {', '.join(missing)}. Use the template.")

    def cell(values, key, limit=254) -> str:
        index = columns.get(key)
        return _text(values[index], limit) if index is not None and index < len(values) else ""

    parsed: list[ParsedRow] = []
    for number, values in enumerate(rows, start=2):  # row 1 is the header
        row = ParsedRow(
            number,
            cell(values, "email").lower(),
            cell(values, "full_name", 150),
            cell(values, "department", 20).upper(),
            cell(values, "role", 20).upper(),
        )
        if not any((row.email, row.full_name, row.department, row.role)):
            continue  # blank line
        parsed.append(row)
        if len(parsed) > settings.IMPORT_MAX_ROWS:
            raise FileRejected(f"The file has more than {settings.IMPORT_MAX_ROWS} employee rows. Split it up.")
    if not parsed:
        raise FileRejected("The workbook has no employee rows.")
    return parsed


def validate_rows(rows: list[ParsedRow], organization) -> list[ParsedRow]:
    """Attach row-level errors in place. Rows without errors can be imported."""
    departments = {d.code: d for d in Department.objects.filter(organization=organization, status="ACTIVE")}
    existing = {
        e.lower()
        for e in User.objects.filter(email__in=[r.email for r in rows if r.email]).values_list("email", flat=True)
    }
    first_row_for_email: dict[str, int] = {}

    for row in rows:

        def problem(field_name, message, row=row):
            row.errors.append({"field": field_name, "message": message})

        if not row.email:
            problem("email", "Email is required.")
        else:
            try:
                validate_email(row.email)
            except ValidationError:
                problem("email", "Not a valid email address.")
            else:
                if row.email in first_row_for_email:
                    problem("email", f"Duplicate of row {first_row_for_email[row.email]}.")
                else:
                    first_row_for_email[row.email] = row.row
                if row.email in existing:
                    problem("email", "A user with this email already exists.")
        if not row.full_name:
            problem("full_name", "Full name is required.")
        elif row.full_name.startswith(FORMULA_PREFIXES):
            problem("full_name", "Names may not start with = + - @ (they would act as spreadsheet formulas).")
        if row.department and row.department not in departments:
            problem("department", f'No active department with code "{row.department}".')
        if row.role not in IMPORTABLE_ROLES:
            if row.role in {"ORG_ADMIN", "SEC_ADMIN"}:
                problem(
                    "role",
                    "Privileged roles cannot be assigned by import. Import as Employee, then change the "
                    "role after the person has activated their account.",
                )
            else:
                problem("role", f'Unknown role "{row.role}". Leave it empty or use EMPLOYEE.')
    return rows
