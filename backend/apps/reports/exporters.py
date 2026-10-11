"""CSV / Excel / PDF writers for report rows.

Spreadsheet formula injection: a cell that starts with = + - @ (or tab / carriage return) can run as a
formula when a spreadsheet opens the file. Every text value written here is neutralised:
  * CSV  -> prefixed with an apostrophe (the OWASP-recommended mitigation)
  * XLSX -> written explicitly as a text cell, never as a formula
  * PDF  -> no formulas, but text is XML-escaped so it cannot inject reportlab markup
"""

import csv
import io
from datetime import UTC, date, datetime
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")


def display(value):
    """A plain, timezone-explicit representation for files."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, datetime):
        return value.astimezone(UTC).strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    return value


def csv_safe(value) -> str:
    text = str(display(value))
    return "'" + text if text.startswith(FORMULA_TRIGGERS) else text


def to_csv(headers: list[str], rows: list[list]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")  # CRLF: what Excel expects
    writer.writerow([csv_safe(h) for h in headers])
    for row in rows:
        writer.writerow([csv_safe(v) for v in row])
    return ("﻿" + buffer.getvalue()).encode("utf-8")  # BOM so Excel detects UTF-8


def to_xlsx(title: str, headers: list[str], rows: list[list]) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title[:31] or "Report")
    widths = [len(h) for h in headers]
    for row in rows[:500]:
        for i, value in enumerate(row):
            widths[i] = max(widths[i], min(len(str(display(value))), 60))
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width + 2

    def cell(value, bold=False):
        value = display(value)
        c = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            c.data_type = "s"  # always text: a value like "=HYPERLINK(...)" must never become a formula
        if bold:
            c.font = Font(bold=True)
        return c

    sheet.append([cell(h, bold=True) for h in headers])
    for row in rows:
        sheet.append([cell(v) for v in row])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_pdf(title: str, subtitle: str, headers: list[str], rows: list[list], truncated_note: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=7, leading=8.5)
    head = styles["BodyText"].clone("head", fontSize=7, leading=8.5, textColor=colors.white)
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title=title,
        author="SecureVault",
    )

    def p(value, style):
        return Paragraph(escape(str(display(value))), style)

    data = [[p(h, head) for h in headers]] + [[p(v, small) for v in row] for row in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story = [
        Paragraph(escape(title), styles["Title"]),
        Paragraph(escape(subtitle), styles["Normal"]),
        Spacer(1, 4 * mm),
    ]
    if truncated_note:
        story += [Paragraph(escape(truncated_note), styles["Italic"]), Spacer(1, 2 * mm)]
    story.append(table)
    document.build(story)
    return buffer.getvalue()
