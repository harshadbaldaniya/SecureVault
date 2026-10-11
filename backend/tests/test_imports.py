import io
import zipfile

import pytest
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from openpyxl import Workbook, load_workbook

from apps.accounts.models import Invitation, User
from apps.audit.models import AuditAction, AuditLog
from apps.departments.models import Department
from apps.imports.models import ImportJob
from apps.imports.services import run_queued
from apps.notifications.models import Notification

pytestmark = pytest.mark.django_db

URL = "/api/v1/imports/"
HEADER = ["email", "full_name", "department", "role"]


def workbook_bytes(rows, header=HEADER) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if header:
        sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def upload(client, data: bytes, name="employees.xlsx"):
    return client.post(URL, {"file": SimpleUploadedFile(name, data)}, format="multipart")


@pytest.fixture
def eng(org):
    return Department.objects.create(organization=org, name="Engineering", code="ENG")


# --------------------------------------------------------------------------- row validation and results


def test_import_reports_successful_and_failed_rows(
    client_for, admin, eng, employee, django_capture_on_commit_callbacks
):
    rows = [
        ["ok.one@example.test", "Ok One", "eng", ""],  # department code is case-insensitive
        ["ok.two@example.test", "Ok Two", "", "EMPLOYEE"],
        ["not-an-email", "Bad Email", "", ""],
        ["no.name@example.test", "", "", ""],
        ["ghost.dept@example.test", "Ghost Dept", "XXX", ""],
        ["wants.admin@example.test", "Wants Admin", "", "SEC_ADMIN"],
        ["ok.one@example.test", "Duplicate", "", ""],
        [employee.email, "Already Here", "", ""],
        ["formula@example.test", '=HYPERLINK("http://evil.example","click")', "", ""],
        ["weird.role@example.test", "Weird Role", "", "SUPERUSER"],
    ]
    with django_capture_on_commit_callbacks(execute=True):
        response = upload(client_for(admin), workbook_bytes(rows))
    assert response.status_code == 201
    job = response.json()
    assert (job["status"], job["total_rows"], job["success_count"], job["failed_count"]) == ("COMPLETED", 10, 2, 8)

    by_row = {e["row"]: e for e in job["errors"]}
    expect = {
        4: "email",
        5: "full_name",
        6: "department",
        7: "role",
        8: "email",
        9: "email",
        10: "full_name",
        11: "role",
    }
    assert {row: by_row[row]["errors"][0]["field"] for row in expect} == expect
    assert "Duplicate of row 2" in by_row[8]["errors"][0]["message"]
    assert "Privileged roles cannot be assigned by import" in by_row[7]["errors"][0]["message"]

    invited = User.objects.filter(email__in=["ok.one@example.test", "ok.two@example.test"])
    assert {u.status for u in invited} == {"INVITED"} and {u.role for u in invited} == {"EMPLOYEE"}
    assert invited.get(email="ok.one@example.test").department == eng
    assert not User.objects.filter(email="wants.admin@example.test").exists()

    assert sorted(m.to[0] for m in mail.outbox) == ["ok.one@example.test", "ok.two@example.test"]
    assert "/accept-invite#token=" in mail.outbox[0].body
    assert "#token=" not in response.content.decode()  # links go to invitees, never to the uploader
    assert ImportJob.objects.get(pk=job["id"]).pending_rows == []
    assert AuditLog.objects.filter(action=AuditAction.EMPLOYEE_IMPORT, metadata__succeeded=2).exists()
    assert Notification.objects.filter(recipient=admin, type="IMPORT_COMPLETED").exists()


def test_retry_imports_fixed_rows_without_creating_duplicates(client_for, admin, org, eng):
    rows = [["first@example.test", "First", "QA", ""], ["second@example.test", "Second", "ENG", ""]]
    client = client_for(admin)
    job = upload(client, workbook_bytes(rows)).json()
    assert (job["success_count"], job["failed_count"]) == (1, 1)

    Department.objects.create(organization=org, name="Quality", code="QA")  # the admin fixes the cause
    retried = client.post(f"{URL}{job['id']}/retry/").json()
    assert retried["retry_of"] == job["id"] and (retried["success_count"], retried["failed_count"]) == (1, 0)
    assert User.objects.filter(email="first@example.test").count() == 1

    again = client.post(f"{URL}{job['id']}/retry/").json()  # retrying the same original again
    assert (again["success_count"], again["failed_count"]) == (0, 1)
    assert "already exists" in again["errors"][0]["errors"][0]["message"]
    assert User.objects.filter(email="first@example.test").count() == 1


def test_large_imports_are_queued_for_the_worker(client_for, admin, settings):
    settings.IMPORT_SYNC_ROW_LIMIT = 1
    rows = [[f"u{i}@example.test", f"User {i}", "", ""] for i in range(3)]
    job = upload(client_for(admin), workbook_bytes(rows)).json()
    assert job["status"] == "QUEUED" and job["success_count"] == 0
    call_command("run_worker", "--once")
    stored = ImportJob.objects.get(pk=job["id"])
    assert (stored.status, stored.success_count) == ("COMPLETED", 3)
    assert run_queued() == 0  # nothing left


# --------------------------------------------------------------------------- hostile and malformed files


def test_hostile_or_malformed_files_are_rejected(client_for, admin, settings):
    client = client_for(admin)
    valid = workbook_bytes([["a@example.test", "A", "", ""]])

    def failed(response):
        assert response.status_code == 201  # the job is recorded for the import history...
        body = response.json()
        assert body["status"] == "FAILED" and body["success_count"] == 0  # ...but nothing is imported
        return body["file_error"]

    assert ".xlsx" in failed(upload(client, valid, "employees.csv"))
    assert "valid .xlsx" in failed(upload(client, b"email,full_name\nx@y.z,X", "fake.xlsx"))

    macro = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(valid)) as src, zipfile.ZipFile(macro, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/vbaProject.bin", b"\x00" * 64)
    assert "macros" in failed(upload(client, macro.getvalue()))

    bomb = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(valid)) as src, zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/media/padding.bin", b"\x00" * (60 * 1024 * 1024))  # compresses to ~60 KB
    assert len(bomb.getvalue()) < settings.IMPORT_MAX_BYTES
    assert "too large when uncompressed" in failed(upload(client, bomb.getvalue()))

    assert "Missing required column" in failed(upload(client, workbook_bytes([["x@y.z"]], header=["email"])))
    assert "no employee rows" in failed(upload(client, workbook_bytes([])))
    settings.IMPORT_MAX_ROWS = 2
    assert "more than 2" in failed(upload(client, workbook_bytes([[f"u{i}@x.test", "U", "", ""] for i in range(3)])))
    assert not User.objects.filter(email__endswith="@x.test").exists()


def test_xml_entity_expansion_is_blocked_by_defusedxml(client_for, admin):
    """A 'billion laughs' style DTD in the worksheet XML must be refused, not expanded, and must not 500."""
    plain = io.BytesIO(workbook_bytes([["a@example.test", "Alice"]], header=["email", "full_name"]))
    sheet_part = "xl/worksheets/sheet1.xml"
    hostile = io.BytesIO()
    with zipfile.ZipFile(plain) as src, zipfile.ZipFile(hostile, "w") as dst:
        for name in src.namelist():
            content = src.read(name)
            if name == sheet_part:
                xml = content.decode()
                dtd = '<!DOCTYPE worksheet [<!ENTITY lol "lol"><!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;">]>'
                xml = xml.replace("<worksheet", dtd + "<worksheet", 1).replace(">Alice<", ">&lol2;<", 1)
                assert "&lol2;" in xml  # the payload really is in place
                content = xml.encode()
            dst.writestr(name, content)
    response = upload(client_for(admin), hostile.getvalue())
    assert response.status_code == 201
    body = response.json()
    assert body["status"] == "FAILED" and body["file_error"] == "The workbook could not be read."
    assert not User.objects.filter(email="a@example.test").exists()


def test_upload_size_limit(client_for, admin, settings):
    settings.IMPORT_MAX_BYTES = 1000
    response = upload(client_for(admin), b"x" * 2000)
    assert response.status_code == 400
    assert not ImportJob.objects.exists()


# --------------------------------------------------------------------------- access and template


def test_only_organization_admins_can_import(client_for, employee, sec_admin):
    data = workbook_bytes([["a@example.test", "A", "", ""]])
    assert upload(client_for(employee), data).status_code == 403
    assert upload(client_for(sec_admin), data).status_code == 403
    assert not ImportJob.objects.exists()


def test_import_history_is_scoped_to_the_organization(client_for, admin, other_admin):
    job_id = upload(client_for(admin), workbook_bytes([["a@example.test", "A", "", ""]])).json()["id"]
    assert client_for(other_admin).get(URL).json()["count"] == 0
    assert client_for(other_admin).get(f"{URL}{job_id}/").status_code == 404


def test_template_is_a_valid_workbook_with_the_expected_columns(client_for, admin):
    response = client_for(admin).get(f"{URL}template/")
    assert response.status_code == 200
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    sheet = load_workbook(io.BytesIO(response.content), read_only=True).worksheets[0]
    assert next(sheet.iter_rows(values_only=True)) == tuple(HEADER)


def test_imported_invitation_can_be_accepted(client_for, admin, django_capture_on_commit_callbacks):
    with django_capture_on_commit_callbacks(execute=True):
        upload(client_for(admin), workbook_bytes([["new.hire@example.test", "New Hire", "", ""]]))
    token = mail.outbox[0].body.split("#token=", 1)[1].split()[0]
    from rest_framework.test import APIClient

    response = APIClient().post(
        "/api/v1/auth/invitations/accept/", {"token": token, "password": "Fresh-Hire-Passphrase-31"}, format="json"
    )
    assert response.status_code == 200
    assert User.objects.get(email="new.hire@example.test").status == "ACTIVE"
    assert Invitation.objects.get(user__email="new.hire@example.test").accepted_at is not None
